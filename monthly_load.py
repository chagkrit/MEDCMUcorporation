"""Read Brand Scan exports (Social Metric / Owned / Earned) into one normalised post table.

Every export workbook has an Overview sheet (Timeframe, Generated At) and one sheet per platform.
Column sets differ per platform and per export family, so values are looked up by header name and
anything absent becomes NaN (missing) - never 0.
"""
from pathlib import Path
import re

import numpy as np
import openpyxl
import pandas as pd

BRANDS = [  # (key, short, label, full Thai name in the files); order == colour slot order
    ("MEDCMU", "MEDCMU", "คณะแพทยศาสตร์ มช.", "คณะแพทยศาสตร์ มหาวิทยาลัยเชียงใหม่"),
    ("SI", "ศิริราช", "คณะแพทยศาสตร์ศิริราชพยาบาล", "คณะแพทยศาสตร์ศิริราชพยาบาล"),
    ("CU", "จุฬาฯ", "โรงพยาบาลจุฬาลงกรณ์ สภากาชาดไทย", "โรงพยาบาลจุฬาลงกรณ์ สภากาชาดไทย"),
    ("SMV", "สมิติเวช", "โรงพยาบาลสมิติเวช", "โรงพยาบาลสมิติเวช"),
]
BRAND_MAP = {b[3]: b[0] for b in BRANDS}
PLATFORMS = ["facebook", "instagram", "tiktok", "youtube", "twitter"]
SHEET_PLATFORM = {"Facebook": "facebook", "Instagram": "instagram", "TikTok": "tiktok", "YouTube": "youtube", "X": "twitter"}
FAMILY_SCOPE = {"social_metric": "owned", "owned": "owned", "earned": "earned"}
REQUIRED_COLS = ("Date", "Brand Name", "Author", "Message", "Post URL", "Engagement")   # a rename in a new export must stop the build
COLS = ["date", "month", "brand", "platform", "scope", "author", "text", "url", "engagement", "views",
        "comment", "share", "family", "src_file", "generated_at", "tf_start", "tf_end"]


def family_of(path):
    n = Path(path).name
    if "Social Metric" in n:
        return "social_metric"
    if " Owned " in n:
        return "owned"
    if " Earned " in n:
        return "earned"
    raise ValueError(f"unrecognised export file name: {n}")


def num(v):
    """Cell -> float; None / '' / 'NaN' / unparsable -> NaN (missing, not zero)."""
    if v is None:
        return np.nan
    if isinstance(v, str):
        s = v.strip().replace(",", "")
        if s == "" or s.lower() == "nan":
            return np.nan
        try:
            return float(s)
        except ValueError:
            return np.nan
    return float(v)


def clean_text(v):
    if v is None:
        return None
    s = str(v).strip()
    return None if s == "" or s.lower() == "nan" else s


def read_overview(wb):
    gen = tf = None
    for r in wb["Overview"].iter_rows(values_only=True):
        cells = [c for c in r if c is not None]
        if len(cells) >= 2 and cells[0] == "Generated At":
            gen = str(cells[1])
        if len(cells) >= 2 and cells[0] == "Timeframe":
            tf = str(cells[1])
    if not gen or not tf:
        raise ValueError("Overview sheet lacks 'Generated At' / 'Timeframe'")
    a, b = re.findall(r"\d{4}-\d{2}-\d{2}", tf)[:2]
    return gen, a, b


def read_file(path):
    fam = family_of(path)
    scope = FAMILY_SCOPE[fam]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    gen, tf0, tf1 = read_overview(wb)
    name = Path(path).name
    out = []
    for ws in wb.worksheets:
        if ws.title == "Overview":
            continue
        if ws.title not in SHEET_PLATFORM:
            raise ValueError(f"unknown platform sheet {ws.title!r} in {name}")
        pf = SHEET_PLATFORM[ws.title]
        it = ws.iter_rows(values_only=True)
        hdr = [str(h) for h in next(it)]
        ix = {h: i for i, h in enumerate(hdr)}
        missing = [c for c in REQUIRED_COLS if c not in ix]
        if missing:
            raise ValueError(f"missing required column(s) {missing} in {name}/{ws.title}")

        def g(row, col):
            i = ix.get(col)
            return row[i] if i is not None and i < len(row) else None

        for row in it:
            if row is None or all(c is None for c in row):
                continue
            brand = BRAND_MAP.get(g(row, "Brand Name"))
            if brand is None:
                raise ValueError(f"unmapped brand {g(row, 'Brand Name')!r} in {name}/{ws.title}")
            date = str(g(row, "Date"))[:10]
            if pf == "twitter":
                comment = num(g(row, "Reply"))
                rp, qt = num(g(row, "Repost")), num(g(row, "Quote"))
                share = np.nan if (np.isnan(rp) and np.isnan(qt)) else float(np.nansum([rp, qt]))
            else:
                comment, share = num(g(row, "Comment")), num(g(row, "Share"))
            out.append((date, date[:7], brand, pf, scope, clean_text(g(row, "Author")), clean_text(g(row, "Message")),
                        clean_text(g(row, "Post URL")), num(g(row, "Engagement")), num(g(row, "View")),
                        comment, share, fam, name, gen, tf0, tf1))
    wb.close()
    return pd.DataFrame(out, columns=COLS)


def load_all(data_dir):
    files = sorted(p for p in Path(data_dir).glob("*.xlsx") if not p.name.startswith("~$"))
    if not files:
        raise FileNotFoundError(f"no .xlsx files in {data_dir}")
    return pd.concat([read_file(p) for p in files], ignore_index=True)
