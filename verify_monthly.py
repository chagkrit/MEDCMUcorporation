#!/usr/bin/env python3
"""Independent recount of monthly_ledger.csv and monthly_data.json from the raw xlsx files.
Shares no code with build_monthly.py. Checks: every ledger metric per cell (rows, posts, blank, engagement, views,
comment, share, mean, median, top1, top3), the source file of each cell, month status/span/source, the daily series
totals, the top-posts lists and the earned top-author lists. A cell whose values are all missing must be missing."""
import calendar
import csv
import glob
import json
import os
import re
import statistics
import sys
from collections import defaultdict
from datetime import date, timedelta

import openpyxl

BRAND = {"คณะแพทยศาสตร์ มหาวิทยาลัยเชียงใหม่": "MEDCMU", "คณะแพทยศาสตร์ศิริราชพยาบาล": "SI",
         "โรงพยาบาลจุฬาลงกรณ์ สภากาชาดไทย": "CU", "โรงพยาบาลสมิติเวช": "SMV"}
PLAT = {"Facebook": "facebook", "Instagram": "instagram", "TikTok": "tiktok", "YouTube": "youtube", "X": "twitter"}
METRICS = ("rows", "posts", "blank", "eng", "views", "comment", "share", "mean", "median", "top1", "top3")
TOL = 1e-6


def fnum(v):
    try:
        s = str(v).strip().replace(",", "")
        return 0.0 if s in ("", "None", "NaN", "nan") else float(s)
    except ValueError:
        return 0.0


def fopt(v):
    """Like fnum but a missing cell stays None (so an all-missing cell must be 'missing', not 0)."""
    try:
        s = str(v).strip().replace(",", "")
        return None if s in ("", "None", "NaN", "nan") else float(s)
    except ValueError:
        return None


def col(r, ix, name):
    i = ix.get(name)
    return r[i] if i is not None and i < len(r) else None


def has_text(v):
    return v is not None and str(v).strip() not in ("", "nan", "NaN")


def read_files(data_dir):
    files = []
    for path in sorted(glob.glob(os.path.join(data_dir, "*.xlsx"))):
        name = os.path.basename(path)
        if name.startswith("~$"):
            continue
        scope = "earned" if " Earned " in name else "owned"
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        gen = tf = None
        for r in wb["Overview"].iter_rows(values_only=True):
            c = [x for x in r if x is not None]
            if len(c) >= 2 and c[0] == "Generated At":
                gen = str(c[1])
            if len(c) >= 2 and c[0] == "Timeframe":
                tf = str(c[1])
        tf0, tf1 = re.findall(r"\d{4}-\d{2}-\d{2}", tf)[:2]
        rows, derived = [], 0
        for ws in wb.worksheets:
            if ws.title == "Overview":
                continue
            start, yt_lc = len(rows), []
            it = ws.iter_rows(values_only=True)
            h = [str(x) for x in next(it)]
            ix = {k: i for i, k in enumerate(h)}
            for r in it:
                if r is None or all(x is None for x in r):
                    continue
                day = str(r[ix["Date"]])[:10]
                if ws.title == "X":                                   # Reply -> comment ; Repost + Quote -> share
                    comment = fopt(col(r, ix, "Reply"))
                    rp, qt = fopt(col(r, ix, "Repost")), fopt(col(r, ix, "Quote"))
                    share = None if rp is None and qt is None else (rp or 0.0) + (qt or 0.0)
                else:
                    comment, share = fopt(col(r, ix, "Comment")), fopt(col(r, ix, "Share"))
                au = col(r, ix, "Author")
                rows.append((day[:7], BRAND[r[ix["Brand Name"]]], PLAT[ws.title], day, r[ix["Message"]], r[ix["Post URL"]],
                             fnum(r[ix["Engagement"]]), fopt(col(r, ix, "View")), comment, share, str(au).strip() if has_text(au) else None))
                if ws.title == "YouTube":
                    yt_lc.append((fopt(col(r, ix, "Like")) or 0.0) + (comment or 0.0))
            # YouTube Engagement column is 0 on every row of the Social Metric sheets; in the Owned export it equals Like + Comment
            if ws.title == "YouTube" and rows[start:] and all(x[6] == 0 for x in rows[start:]) and any(v > 0 for v in yt_lc):
                for i, v in enumerate(yt_lc):
                    rows[start + i] = rows[start + i][:6] + (v,) + rows[start + i][7:]
                derived += len(yt_lc)
        wb.close()
        files.append(dict(name=name, scope=scope, gen=gen, tf0=tf0, tf1=tf1, rows=rows, derived=derived))
    return files


def month_list(months):
    y, m = map(int, min(months).split("-"))
    y2, m2 = map(int, max(months).split("-"))
    out = []
    while (y, m) <= (y2, m2):
        out.append(f"{y}-{m:02d}")
        m += 1
        if m == 13:
            y, m = y + 1, 1
    return out


def span_status(f, mo):
    d0 = f"{mo}-01"
    y, m = map(int, mo.split("-"))
    d1 = f"{mo}-{calendar.monthrange(y, m)[1]:02d}"
    prev = (date.fromisoformat(f["gen"][:10]) - timedelta(days=1)).isoformat()   # an export cannot hold a complete day it was generated on
    a, b = max(f["tf0"], d0), min(f["tf1"], prev, d1)
    return ("complete" if (a == d0 and b == d1) else "partial"), [a, b]


def main(data_dir, ledger_path, data_path):
    files = read_files(data_dir)
    newest = {}                                    # (month, scope) -> file with newest Generated At
    for f in files:
        for mo in {r[0] for r in f["rows"]}:
            k = (mo, f["scope"])
            if k not in newest or (f["gen"], f["name"]) > (newest[k]["gen"], newest[k]["name"]):
                newest[k] = f
    agg = defaultdict(lambda: {"rows": 0, "posts": 0, "eng": 0.0, "views": None, "comment": None, "share": None, "texteng": [], "auth": defaultdict(float)})
    seen = set()
    for f in files:
        for mo, b, p, day, msg, url, eng, views, comment, share, author in f["rows"]:
            if newest[(mo, f["scope"])] is not f:
                continue
            if url is not None and str(url).strip() not in ("", "nan"):
                dk = (b, p, f["scope"], str(url).strip(), day, eng)
                if dk in seen:
                    continue
                seen.add(dk)
            a = agg[(mo, b, p, f["scope"])]
            a["rows"] += 1
            if has_text(msg):
                a["posts"] += 1
                a["texteng"].append(eng)
            a["eng"] += eng
            for name, v in (("views", views), ("comment", comment), ("share", share)):
                if v is not None:
                    a[name] = (a[name] or 0.0) + v        # stays None when every row in the cell is missing
            if f["scope"] == "earned" and author is not None:
                a["auth"][author] += eng
    for a in agg.values():                                # concentration/mean/median are over text-bearing posts
        t = sorted(a["texteng"], reverse=True)
        tt = sum(t)
        a["blank"] = a["rows"] - a["posts"]
        a["mean"] = round(tt / len(t), 2) if t else None
        a["median"] = round(statistics.median(t), 2) if t else None
        a["top1"] = round(t[0] / tt * 100, 1) if t and tt > 0 else None
        a["top3"] = round(sum(t[:3]) / tt * 100, 1) if t and tt > 0 else None
    want, cells, bad = {}, set(), []
    for r in csv.DictReader(open(ledger_path, encoding="utf-8-sig")):
        key = (r["month"], r["brand"], r["platform"], r["scope"])
        cells.add(key)
        if key in agg and r["source"] != newest[(key[0], key[3])]["name"]:      # every ledger row carries the source file
            bad.append((key, "source of " + r["metric"], newest[(key[0], key[3])]["name"], r["source"]))
        if r["metric"] in METRICS:
            want[key + (r["metric"],)] = float(r["value"])

    def same(exp, got):
        return (exp is None) == (got is None) and (exp is None or abs(got - exp) <= TOL)

    for k, a in agg.items():
        for metric in METRICS:
            exp, got = a[metric], want.get(k + (metric,))      # the ledger omits a metric that is missing
            if not same(exp, got):
                bad.append((k, metric, exp, got))
    if cells != set(agg):
        bad.append(("cell-set differs", sorted(cells ^ set(agg))[:5], None, None))

    data = json.load(open(data_path, encoding="utf-8"))
    allm = {r[0] for f in files for r in f["rows"]}
    if [m["mo"] for m in data["months"]] != month_list(allm):
        bad.append(("month list differs", [m["mo"] for m in data["months"]], month_list(allm), None))
    for m in data["months"]:
        for sc in ("owned", "earned"):
            f = newest.get((m["mo"], sc))
            exp = ("nodata", None, None) if f is None else (*span_status(f, m["mo"]), f["name"])
            got = (m[sc]["status"], m[sc]["span"], m[sc]["src"])
            if exp != got:
                bad.append((m["mo"], sc + " status/span/src", exp, got))
    der_exp = {f["name"]: f["derived"] for f in files if f["derived"]}
    der_got = {d["file"]: d["rows"] for d in data["meta"]["derived"]}
    if der_exp != der_got:
        bad.append(("derived-engagement files differ", der_exp, der_got, None))
    def norm(msg):
        return re.sub(r"\s+", "", str(msg).strip().lower())[:80]
    ov_exp, cand = defaultdict(lambda: [0, 0, 0, 0.0, 0.0]), defaultdict(list)
    for f in files:
        for mo in {r[0] for r in f["rows"]}:
            cand[(mo, f["scope"])].append(f)
    for (mo, sc), fl in cand.items():
        if len(fl) < 2:
            continue
        fl = sorted(fl, key=lambda f: (f["gen"], f["name"]), reverse=True)
        wrows = [r for r in fl[0]["rows"] if r[0] == mo]
        for other in fl[1:]:
            orows = [r for r in other["rows"] if r[0] == mo]
            d0, d1 = min(r[3] for r in orows), max(r[3] for r in orows)
            wwin = [r for r in wrows if d0 <= r[3] <= d1]
            for br in {r[1] for r in wwin + orows}:
                for pf in {r[2] for r in wwin + orows if r[1] == br}:
                    ga = [r for r in wwin if r[1] == br and r[2] == pf]
                    gb = [r for r in orows if r[1] == br and r[2] == pf]
                    ka = {r[3] + "|" + norm(r[4]) for r in ga if has_text(r[4])}
                    kb = {r[3] + "|" + norm(r[4]) for r in gb if has_text(r[4])}
                    e = ov_exp[(mo, sc, br)]
                    e[0] += sum(1 for r in ga if has_text(r[4])); e[1] += sum(1 for r in gb if has_text(r[4])); e[2] += len(ka & kb)
                    e[3] += sum(r[6] for r in ga); e[4] += sum(r[6] for r in gb)
    ov_got = defaultdict(lambda: [0, 0, 0, 0.0, 0.0])
    for r in data["quality"]["overlap"]:
        g = ov_got[(r["month"], r["scope"], r["brand"])]
        g[0] += r["win_text"]; g[1] += r["other_text"]; g[2] += r["matched"]; g[3] += r["win_eng"]; g[4] += r["other_eng"]
    if set(ov_exp) != set(ov_got) or any(any(abs(x - y) > TOL for x, y in zip(ov_exp[k], ov_got[k])) for k in ov_exp if k in ov_got):
        bad.append(("overlap table differs", len(ov_exp), len(ov_got), None))
    series = defaultdict(lambda: [0.0, 0])               # (brand, platform, scope) -> engagement, text posts over all months
    for (mo, b, p, s), a in agg.items():
        series[(b, p, s)][0] += a["eng"]
        series[(b, p, s)][1] += a["posts"]
    if {tuple(k.split("|")) for k in data["daily"]} != set(series):
        bad.append(("daily series set differs", None, None, None))
    for k, v in data["daily"].items():
        key = tuple(k.split("|"))
        if key in series:
            e, n = sum(x for x in v["eng"] if x is not None), sum(x for x in v["posts"] if x is not None)
            if abs(e - series[key][0]) > TOL or n != series[key][1]:
                bad.append(("daily " + k, (series[key][0], series[key][1]), (e, n), None))
    top_got = defaultdict(list)
    for r in data["top"]:
        top_got[(r["m"], r["b"], r["p"], r["s"])].append(r["e"])
    top_exp = {k: [int(x) for x in sorted(a["texteng"], reverse=True)[:10]] for k, a in agg.items() if a["texteng"]}
    if {k: sorted(v, reverse=True) for k, v in top_got.items()} != top_exp:
        bad.append(("top-posts lists differ", len(top_got), len(top_exp), None))
    au_got = defaultdict(list)
    for r in data["authors"]:
        au_got[(r["m"], r["b"], r["p"])].append(r["e"])
    au_exp = {(k[0], k[1], k[2]): sorted(a["auth"].values(), reverse=True)[:10] for k, a in agg.items() if k[3] == "earned" and a["auth"]}
    if set(au_got) != set(au_exp) or any(len(au_got[k]) != len(au_exp[k]) or any(abs(x - y) > TOL for x, y in zip(sorted(au_got[k], reverse=True), au_exp[k])) for k in au_exp if k in au_got):
        bad.append(("author lists differ", len(au_got), len(au_exp), None))

    print(f"recounted cells: {len(agg)}; ledger cells: {len(cells)}; ledger values compared: {len(agg) * len(METRICS)} (+ every ledger row's source label); "
          f"months: {len(data['months'])}; daily series: {len(data['daily'])}; top-post groups: {len(top_exp)}; author groups: {len(au_exp)}; overlap groups: {len(ov_exp)}; mismatches: {len(bad)}")
    for b in bad[:20]:
        print("  MISMATCH", b)
    return 1 if bad else 0


if __name__ == "__main__":
    here = os.path.dirname(os.path.abspath(__file__))
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(here, "..", "ZocialEye"),
                  sys.argv[2] if len(sys.argv) > 2 else os.path.join(here, "monthly_ledger.csv"),
                  sys.argv[3] if len(sys.argv) > 3 else os.path.join(here, "monthly_data.json")))
